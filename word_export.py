from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY = RGBColor(17, 36, 58)
BLUE = RGBColor(20, 104, 160)
MUTED = RGBColor(100, 116, 139)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def _add_labeled(paragraph, label: str, value: str) -> None:
    run = paragraph.add_run(label)
    run.bold = True
    run.font.color.rgb = NAVY
    paragraph.add_run(value or "—")
    paragraph.paragraph_format.space_after = Pt(4)


def build_daily_pack(
    drafts: list[dict[str, Any]],
    employers_by_id: dict[str, dict[str, Any]],
    pack_day: date,
) -> bytes:
    """Build an in-memory DOCX containing the supplied email drafts; it never sends email."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title_style = document.styles["Title"]
    title_style.font.name = "Georgia"
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = NAVY
    title_style.paragraph_format.space_after = Pt(5)

    heading = document.styles["Heading 1"]
    heading.font.name = "Georgia"
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = BLUE
    heading.paragraph_format.space_before = Pt(4)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True

    document.core_properties.title = f"Air University internship email pack — {pack_day.isoformat()}"
    document.core_properties.subject = "Reviewed email drafts for manual sending"
    document.core_properties.author = "Air University H-11 Internship Desk"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("AIR UNIVERSITY  ·  H-11 CAMPUS  ·  SUMMER 2027")
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Daily email pack  ·  Manual sending only  ·  Page ")
    _add_page_field(footer)

    document.add_paragraph("Summer 2027 internship outreach", style="Title")
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(f"Daily email pack  ·  {pack_day.strftime('%A, %d %B %Y')}")
    subtitle_run.bold = True
    subtitle_run.font.color.rgb = BLUE
    subtitle.paragraph_format.space_after = Pt(9)
    notice = document.add_paragraph()
    notice.add_run(
        "REVIEWED DRAFTS — No email was sent by this app. Copy each message into your own mail compose window, "
        "confirm the recipient and content, add your usual signature if appropriate, and send manually."
    )
    notice.paragraph_format.space_after = Pt(14)
    for run in notice.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    if not drafts:
        document.add_paragraph("No reviewed email drafts are included in this pack.")

    for index, draft in enumerate(drafts):
        if index:
            document.add_page_break()
        employer = employers_by_id.get(str(draft.get("employer_id")), {})
        company = str(employer.get("company_name") or "Organization")
        contact = " ".join(
            part for part in [employer.get("first_name"), employer.get("last_name")] if part
        ).strip()
        document.add_heading(company, level=1)
        _add_labeled(document.add_paragraph(), "To: ", str(draft.get("to_email") or ""))
        if contact:
            _add_labeled(document.add_paragraph(), "Contact: ", contact)
        _add_labeled(
            document.add_paragraph(),
            "Suggested faculty areas: ",
            ", ".join(draft.get("program_fit") or employer.get("program_fit") or []) or "Review manually",
        )
        subject = document.add_paragraph()
        _add_labeled(subject, "Subject: ", str(draft.get("subject") or ""))
        body = str(draft.get("body_text") or "").strip()
        paragraphs = [part for part in body.split("\n\n") if part.strip()]
        if not paragraphs and body:
            paragraphs = [body]
        for block in paragraphs:
            paragraph = document.add_paragraph()
            lines = block.splitlines()
            for line_index, line in enumerate(lines):
                if line_index:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)
        status_note = document.add_paragraph()
        status_note.paragraph_format.space_before = Pt(10)
        status_run = status_note.add_run("This message is a draft for manual sending; it has not been sent by the app.")
        status_run.italic = True
        status_run.font.size = Pt(8.5)
        status_run.font.color.rgb = MUTED

    output = BytesIO()
    document.save(output)
    return output.getvalue()
